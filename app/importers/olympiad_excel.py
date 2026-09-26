from pathlib import Path

from openpyxl import load_workbook


COLUMN_MAPPING = {
    "Название для бд": "external_id",
    "Название олимпиады": "name",
    "Профиль": "subject",
    "Классы": "grades",
    "Уровень (РСОШ)": "level",
    "Статус": "status",
    "Организатор / Вуз": "organizer",
    "Ссылка на сайт": "website",

    "Регистрация на отборочный этап": "qualification_registration",
    "Проведение отборочного этапа 1": "qualification_stage_1",
    "Проведение отборочного этапа 2": "qualification_stage_2",
    "Проведение отборочного этапа 3": "qualification_stage_3",

    "Регистрация на финал": "final_registration",
    "Проведение финала": "final_stage",

    "Примечание": "notes",
}


REQUIRED_COLUMNS = {
    "Название для бд",
    "Название олимпиады",
    "Профиль",
    "Классы",
}


class OlympiadExcelImporter:

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

    def read(self) -> list[dict]:

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"Файл не найден: {self.file_path}"
            )

        workbook = load_workbook(
            self.file_path,
            data_only=True
        )

        worksheet = workbook.active

        excel_headers = [
            cell.value
            for cell in worksheet[1]
        ]

        missing_columns = REQUIRED_COLUMNS - set(excel_headers)

        if missing_columns:
            raise ValueError(
                "В Excel отсутствуют обязательные колонки: "
                + ", ".join(sorted(missing_columns))
            )

        headers = [
            COLUMN_MAPPING.get(header, header)
            for header in excel_headers
        ]

        rows = []

        for values in worksheet.iter_rows(
            min_row=2,
            values_only=True
        ):
            if not any(value is not None for value in values):
                continue

            row = dict(zip(headers, values))

            rows.append(row)

        return rows